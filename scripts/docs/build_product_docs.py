"""Build the two distributable AIA technical documents from Markdown."""

from __future__ import annotations

from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
FONT = "Malgun Gothic"
ACCENT = "1F4E79"
BORDER = "D9E2F3"
HEADER_FILL = "1F4E79"


def _set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_borders(cell, color: str = BORDER) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def _set_run_font(run, size: float = 10.5, bold: bool | None = None, color: str = "000000") -> None:
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def _remove_paragraph_borders(element) -> None:
    properties = element.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is not None:
        properties.remove(borders)


def _style_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
        _remove_paragraph_borders(style._element)


def _add_text(document: Document, text: str, style: str | None = None) -> None:
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.keep_together = False
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            _set_run_font(run, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            _set_run_font(run, size=9.5, color="404040")
        else:
            run = paragraph.add_run(part)
            _set_run_font(run)


def _add_code(document: Document, lines: list[str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.2)
    paragraph.paragraph_format.right_indent = Inches(0.2)
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(8)
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        _set_run_font(run, size=9, color="203040")
        if index < len(lines) - 1:
            run.add_break()


def _add_table(document: Document, rows: list[list[str]]) -> None:
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_borders(cell)
            if row_index == 0:
                _set_cell_shading(cell, HEADER_FILL)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(value)
            _set_run_font(run, size=9.2, bold=row_index == 0, color="FFFFFF" if row_index == 0 else "000000")
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def _clean_heading(text: str) -> str:
    return re.sub(r"^\d+[.)]\s*", "", text).strip()


def build(markdown_path: Path, output_path: Path) -> None:
    document = Document()
    _style_document(document)
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    index = 0
    in_code = False
    code_lines: list[str] = []
    while index < len(lines):
        line = lines[index]
        if line.startswith("```"):
            if in_code:
                _add_code(document, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not line.strip():
            index += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].startswith("|"):
                table_lines.append(lines[index])
                index += 1
            table_rows = []
            for table_line in table_lines:
                values = [item.strip() for item in table_line.strip().strip("|").split("|")]
                if all(set(item) <= {"-", ":", " "} for item in values):
                    continue
                table_rows.append(values)
            if table_rows:
                _add_table(document, table_rows)
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            text = _clean_heading(heading.group(2))
            style = "Title" if level == 1 and not document.paragraphs else f"Heading {level}"
            paragraph = document.add_paragraph(style=style)
            paragraph.paragraph_format.keep_with_next = True
            _remove_paragraph_borders(paragraph._p)
            run = paragraph.add_run(text)
            _set_run_font(run, size={1: 22, 2: 16, 3: 12}[level], bold=True)
            index += 1
            continue
        bullet = re.match(r"^[-*]\s+(.+)$", line)
        numbered = re.match(r"^\d+\.\s+(.+)$", line)
        if bullet or numbered:
            paragraph = document.add_paragraph(style="List Bullet" if bullet else "List Number")
            _set_run_font(paragraph.add_run((bullet or numbered).group(1)))
            index += 1
            continue
        if line.startswith("> "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            run = paragraph.add_run(line[2:])
            _set_run_font(run, bold=True, color=ACCENT)
            index += 1
            continue
        _add_text(document, line)
        index += 1
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> int:
    output = ROOT / "docs" / "generated"
    build(
        ROOT / "docs" / "technical" / "AIA-RUNTIME-SECURITY-WHITEPAPER.md",
        output / "AIA-Runtime-Security-Technical-Paper.docx",
    )
    build(
        ROOT / "docs" / "wikidocs" / "AIA-사용자-가이드.md",
        output / "AIA-Protected-Run-사용자-가이드.docx",
    )
    print(f"DOCX_OUTPUT={output.resolve()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
